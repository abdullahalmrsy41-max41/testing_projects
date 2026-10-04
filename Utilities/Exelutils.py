import openpyxl
from selenium import webdriver
def getRowCount(filePath, sheetName):
    workbook = openpyxl.load_workbook(filePath)
    sheet = workbook[sheetName]
    return sheet.max_row
def getColumnCount(filePath, sheetName):
    workbook = openpyxl.load_workbook(filePath)
    sheet = workbook[sheetName]
    return sheet.max_column
def readData(filePath, sheetName, rowCount, colCount):
    workbook = openpyxl.load_workbook(filePath)
    sheet = workbook[sheetName]
    return sheet.cell(row=rowCount, column=colCount).value
def writeData(filePath, sheetName, row_num, col_num,data):
    workbook = openpyxl.load_workbook(filePath)
    sheet = workbook[sheetName]
    sheet.cell(row=row_num, column=col_num).value=data
    workbook.save(filePath)